"""docread：超时（协作式截止时间 + 线程等待超时，不阻塞事件循环）、文件大小、zip 解压总量/条目数/压缩比上限。"""
import asyncio
import time
import zipfile

import pytest
from conftest import run

from ee.core import docread as D
from ee.tools import ToolContext, ToolError, ToolState
from ee.tools.read import ReadTool


def _zip(path, entries):
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in entries:
            z.writestr(name, data)


def test_zip_bomb_ratio_rejected(tmp_path):
    p = tmp_path / "bomb.docx"
    _zip(p, [("[Content_Types].xml", "<x/>"), ("word/document.xml", b"\0" * (8 * 1024 * 1024))])
    with pytest.raises(D.DocError, match="压缩炸弹|解压"):
        D.extract(p)


def test_zip_total_size_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(D, "MAX_UNZIPPED", 1024 * 1024)
    import os
    p = tmp_path / "big.xlsx"
    _zip(p, [("a.xml", os.urandom(700 * 1024)), ("b.xml", os.urandom(700 * 1024))])
    with pytest.raises(D.DocError, match="解压后"):
        D.extract(p)


def test_zip_entry_count_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(D, "MAX_ZIP_ENTRIES", 50)
    p = tmp_path / "many.pptx"
    _zip(p, [(f"f{i}.xml", "x") for i in range(80)])
    with pytest.raises(D.DocError, match="条目"):
        D.extract(p)


def test_file_size_limit(tmp_path, monkeypatch):
    monkeypatch.setattr(D, "MAX_DOC_BYTES", 1000)
    p = tmp_path / "a.pdf"
    p.write_bytes(b"%PDF-1.4\n" + b"0" * 2000)
    with pytest.raises(D.DocError, match="太大"):
        D.extract(p)


def test_cooperative_deadline_xlsx(tmp_path, monkeypatch):
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    for i in range(300):
        ws.append([i, "x"])
    p = tmp_path / "t.xlsx"
    wb.save(p)
    monkeypatch.setattr(D, "DOC_TIMEOUT", -1)       # 截止时间已过：解析循环里第一次检查就停
    with pytest.raises(D.DocError, match="超时"):
        D.extract(p)


def test_extract_async_timeout_does_not_block_loop(tmp_path, monkeypatch):
    p = tmp_path / "slow.pdf"
    p.write_bytes(b"%PDF-1.4\n")

    def slow(path, pages=None, sheet=None, budget=None):
        time.sleep(1.5)
        return "late", {}

    monkeypatch.setitem(D._PARSERS, "pdf", slow)
    ticks = []

    async def main():
        async def ticker():
            for _ in range(5):
                ticks.append(time.monotonic())
                await asyncio.sleep(0.05)
        t0 = time.monotonic()
        tk = asyncio.create_task(ticker())
        with pytest.raises(D.DocError, match="超时"):
            await D.extract_async(p, timeout=0.3)
        await tk
        return time.monotonic() - t0

    took = asyncio.run(main())
    assert took < 1.2 and len(ticks) == 5


def test_read_tool_uses_timeout(tmp_path, monkeypatch):
    p = tmp_path / "slow.docx"
    p.write_bytes(b"PK")

    def slow(path, pages=None, sheet=None, budget=None):
        time.sleep(1.0)
        return "late", {}

    monkeypatch.setitem(D._PARSERS, "docx", slow)
    monkeypatch.setattr(D, "DOC_TIMEOUT", 0.2)
    monkeypatch.setattr(D, "_zip_guard", lambda path: None)
    ctx = ToolContext(cwd=tmp_path, state=ToolState(), output_dir=tmp_path / ".out")
    with pytest.raises(ToolError, match="超时"):
        run(ReadTool().run({"path": "slow.docx"}, ctx))


def test_output_chars_capped(tmp_path, monkeypatch):
    import docx
    d = docx.Document()
    for i in range(200):
        d.add_paragraph("段落" * 50)
    p = tmp_path / "long.docx"
    d.save(p)
    monkeypatch.setattr(D, "MAX_OUT_CHARS", 2000)
    text, info = D.extract(p)
    assert len(text) < 4000 and "截断" in text
