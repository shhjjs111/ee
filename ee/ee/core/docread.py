"""PDF / Word / Excel / PowerPoint 转文字（D4-7），给 read 工具和文件预览用。

- PDF（pypdf）：按页抽文字，每页前加「=== 第 N 页 ===」；pages 选页（"3"、"2-5"、"1,3,7-9"），默认前 50 页。
  没有文字层的页（扫描件）会标出来。
- Word .docx（python-docx）：按文档顺序输出段落和表格；标题样式转成 markdown #，列表项加 "- "，表格转 markdown 表格。
- Excel .xlsx/.xlsm（openpyxl，只读、取公式的缓存值）：先列出所有工作表和大小，再输出选定工作表（默认第一个）的
  单元格，每行「行号 | A | B …」，最多 MAX_ROWS 行；sheet 可按名字或序号选。
- PowerPoint .pptx（python-pptx，可选依赖）：每张幻灯片的文字和备注。
结果按 (路径, mtime, 大小, 参数) 缓存，同一文件反复分段读不重复解析。
"""
from __future__ import annotations

import datetime
import re
from collections import OrderedDict
from pathlib import Path

DOC_EXTS = {".pdf": "pdf", ".docx": "docx", ".xlsx": "xlsx", ".xlsm": "xlsx", ".pptx": "pptx"}
OLD_FORMATS = {".doc": "Word 97-2003（.doc）", ".xls": "Excel 97-2003（.xls）", ".ppt": "PowerPoint 97-2003（.ppt）"}
DEFAULT_PDF_PAGES = 50
MAX_ROWS = 2000
MAX_COLS = 60
MAX_CELL = 300

_CACHE: "OrderedDict[tuple, tuple[str, dict]]" = OrderedDict()
_CACHE_MAX = 16


class DocError(Exception):
    pass


def doc_kind(path: Path) -> str | None:
    return DOC_EXTS.get(path.suffix.lower())


def extract(path: Path, pages: str | None = None, sheet: str | int | None = None) -> tuple[str, dict]:
    """返回 (文字, 信息)。信息含 kind、总页数/工作表列表等。"""
    ext = path.suffix.lower()
    if ext in OLD_FORMATS:
        raise DocError(f"不支持旧版 {OLD_FORMATS[ext]} 格式，请让用户另存为新格式（.docx/.xlsx/.pptx）后再读，"
                       "或在 shell 里用已安装的软件转换。")
    kind = DOC_EXTS.get(ext)
    if kind is None:
        raise DocError(f"不是支持的文档格式：{ext}")
    st = path.stat()
    key = (str(path), st.st_mtime_ns, st.st_size, str(pages or ""), str(sheet if sheet is not None else ""))
    hit = _CACHE.get(key)
    if hit is not None:
        _CACHE.move_to_end(key)
        return hit
    fn = {"pdf": _pdf, "docx": _docx, "xlsx": _xlsx, "pptx": _pptx}[kind]
    try:
        text, info = fn(path, pages=pages, sheet=sheet)
    except DocError:
        raise
    except ImportError as e:
        raise DocError(f"缺少解析 {kind} 的组件：{e}") from e
    except Exception as e:  # noqa: BLE001  文件损坏、加密等
        raise DocError(f"无法解析这个 {kind.upper()} 文件：{type(e).__name__}: {e}") from e
    info["kind"] = kind
    _CACHE[key] = (text, info)
    while len(_CACHE) > _CACHE_MAX:
        _CACHE.popitem(last=False)
    return text, info


# ---------------------------------------------------------------- PDF
def parse_pages(spec: str | None, total: int, default: int = DEFAULT_PDF_PAGES) -> list[int]:
    """'3' / '2-5' / '1,3,7-9' / '10-'（到末尾）→ 1 起的页码列表（去重、按序、限制在 1..total）。"""
    if not spec or not str(spec).strip():
        return list(range(1, min(total, default) + 1))
    out: list[int] = []
    for part in re.split(r"[,，\s]+", str(spec).strip()):
        if not part:
            continue
        m = re.fullmatch(r"(\d+)?\s*[-~～]\s*(\d+)?", part)
        if m:
            a = int(m.group(1) or 1)
            b = int(m.group(2) or total)
        elif part.isdigit():
            a = b = int(part)
        else:
            raise DocError(f"pages 格式不对：{spec}（例：3、2-5、1,3,7-9）")
        if a > b:
            a, b = b, a
        out.extend(range(max(1, a), min(total, b) + 1))
    seen: set[int] = set()
    res = [p for p in out if not (p in seen or seen.add(p))]
    if not res:
        raise DocError(f"pages={spec} 超出范围（共 {total} 页）")
    return res


def _pdf(path: Path, pages=None, sheet=None) -> tuple[str, dict]:
    from pypdf import PdfReader
    reader = PdfReader(str(path))
    if reader.is_encrypted:
        try:
            if not reader.decrypt(""):
                raise DocError("PDF 有密码保护，无法读取")
        except DocError:
            raise
        except Exception as e:  # noqa: BLE001
            raise DocError(f"PDF 有密码保护，无法读取（{e}）") from e
    total = len(reader.pages)
    if total == 0:
        return "", {"pages": 0, "shown": []}
    sel = parse_pages(pages, total)
    parts: list[str] = []
    empty: list[int] = []
    for n in sel:
        try:
            t = reader.pages[n - 1].extract_text() or ""
        except Exception as e:  # noqa: BLE001
            t = f"（本页解析失败：{type(e).__name__}）"
        t = _tidy(t)
        if not t.strip():
            empty.append(n)
            t = "（本页没有可提取的文字：可能是扫描图片）"
        parts.append(f"=== 第 {n} 页 ===\n{t}")
    info = {"pages": total, "shown": sel, "empty_pages": empty}
    meta = reader.metadata or {}
    title = getattr(meta, "title", None) if meta else None
    if title:
        info["title"] = str(title)
    return "\n\n".join(parts), info


def _tidy(t: str) -> str:
    t = t.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    t = re.sub(r"[ \t]+\n", "\n", t)
    return re.sub(r"\n{3,}", "\n\n", t).strip("\n")


# ---------------------------------------------------------------- Word
def _docx(path: Path, pages=None, sheet=None) -> tuple[str, dict]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    d = docx.Document(str(path))
    out: list[str] = []
    n_tables = 0
    body = d.element.body
    for child in body.iterchildren():
        tag = child.tag.rsplit("}", 1)[-1]
        if tag == "p":
            p = Paragraph(child, d)
            line = _para_md(p)
            if line is not None:
                out.append(line)
        elif tag == "tbl":
            n_tables += 1
            out.append(_table_md(Table(child, d)))
    text = "\n".join(out)
    text = re.sub(r"\n{3,}", "\n\n", text).strip("\n")
    return text, {"tables": n_tables, "paragraphs": len(d.paragraphs)}


def _para_md(p) -> str | None:
    text = p.text.replace("\r", "").strip()
    style = (p.style.name if p.style is not None else "") or ""
    if not text:
        return ""
    m = re.match(r"(?:Heading|标题)\s*(\d)", style)
    if m:
        return "#" * min(6, int(m.group(1))) + " " + text
    if style == "Title" or style == "标题":
        return "# " + text
    if "List" in style or "列表" in style:
        return "- " + text
    return text


def _table_md(t) -> str:
    rows: list[list[str]] = []
    for r in t.rows:
        cells = []
        prev = None
        for c in r.cells:
            if prev is not None and c._tc is prev:      # 横向合并的单元格只出一次
                cells.append("")
                continue
            prev = c._tc
            cells.append(_cell(c.text))
        rows.append(cells)
    if not rows:
        return ""
    width = max(len(r) for r in rows)
    rows = [r + [""] * (width - len(r)) for r in rows]
    lines = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * width]
    lines += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n" + "\n".join(lines) + "\n"


def _cell(v) -> str:
    s = str(v).replace("\r", "").replace("\n", " ").replace("|", "\\|").strip()
    return s if len(s) <= MAX_CELL else s[:MAX_CELL] + "…"


# ---------------------------------------------------------------- Excel
def _col(n: int) -> str:
    s = ""
    while n:
        n, r = divmod(n - 1, 26)
        s = chr(65 + r) + s
    return s


def _fmt(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(round(v, 10))
    if isinstance(v, datetime.datetime):
        return v.isoformat(sep=" ").removesuffix(" 00:00:00")
    if isinstance(v, (datetime.date, datetime.time, datetime.timedelta)):
        return str(v)
    return _cell(v)


def _xlsx(path: Path, pages=None, sheet=None) -> tuple[str, dict]:
    import openpyxl
    wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    try:
        names = list(wb.sheetnames)
        if not names:
            return "（没有工作表）", {"sheets": []}
        ws = None
        if sheet is None or str(sheet).strip() == "":
            ws = wb[names[0]]
        else:
            s = str(sheet).strip()
            if s in names:
                ws = wb[s]
            elif s.isdigit() and 1 <= int(s) <= len(names):
                ws = wb[names[int(s) - 1]]
            else:
                low = {n.lower(): n for n in names}
                if s.lower() in low:
                    ws = wb[low[s.lower()]]
            if ws is None:
                raise DocError(f"没有工作表 {s}；现有：{', '.join(names)}")
        sheets = []
        for n in names:
            w = wb[n]
            try:
                dim = w.calculate_dimension()
            except Exception:  # noqa: BLE001
                dim = "?"
            sheets.append({"name": n, "dimension": dim, "rows": w.max_row, "cols": w.max_column})
        header = "工作表：" + "；".join(f"{i}. {s['name']}（{s['dimension']}）" for i, s in enumerate(sheets, 1))
        rows_out: list[str] = []
        shown = 0
        more = 0
        max_c = 0
        for r_idx, row in enumerate(ws.iter_rows(values_only=True), start=1):
            vals = [_fmt(v) for v in row[:MAX_COLS]]
            while vals and vals[-1] == "":
                vals.pop()
            if not vals:
                continue
            if shown >= MAX_ROWS:
                more += 1
                continue
            max_c = max(max_c, len(vals))
            rows_out.append((r_idx, vals))
            shown += 1
        lines = [header, f"\n=== 工作表「{ws.title}」===", ""]
        if not rows_out:
            lines.append("（空表）")
        else:
            lines.append("行 | " + " | ".join(_col(i) for i in range(1, max_c + 1)))
            for r_idx, vals in rows_out:
                lines.append(f"{r_idx} | " + " | ".join(vals + [""] * (max_c - len(vals))))
        if more:
            lines.append(f"\n（还有 {more} 行未显示：只显示前 {MAX_ROWS} 个非空行。需要统计/筛选大表时用 shell 跑 Python（openpyxl / pandas）处理）")
        if (ws.max_column or 0) > MAX_COLS:
            lines.append(f"（只显示前 {MAX_COLS} 列，共 {ws.max_column} 列）")
        return "\n".join(lines), {"sheets": sheets, "sheet": ws.title}
    finally:
        wb.close()


# ---------------------------------------------------------------- PowerPoint
def _pptx(path: Path, pages=None, sheet=None) -> tuple[str, dict]:
    try:
        from pptx import Presentation
    except ImportError as e:
        raise DocError("没有安装 python-pptx，暂时不能读 .pptx") from e
    prs = Presentation(str(path))
    total = len(prs.slides)
    sel = set(parse_pages(pages, total, default=total)) if total else set()
    parts = []
    for i, slide in enumerate(prs.slides, start=1):
        if i not in sel:
            continue
        texts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.text_frame.text.strip():
                texts.append(_tidy(shape.text_frame.text))
            elif getattr(shape, "has_table", False):
                rows = [[_cell(c.text) for c in r.cells] for r in shape.table.rows]
                if rows:
                    texts.append("\n".join("| " + " | ".join(r) + " |" for r in rows))
        notes = ""
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame is not None:
            notes = _tidy(slide.notes_slide.notes_text_frame.text)
        body = "\n".join(texts) or "（本页没有文字）"
        if notes:
            body += f"\n[备注] {notes}"
        parts.append(f"=== 第 {i} 张幻灯片 ===\n{body}")
    return "\n\n".join(parts), {"slides": total, "shown": sorted(sel)}
